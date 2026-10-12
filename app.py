import json
import difflib
import io
import os
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path
from flask import (
    Flask, render_template, jsonify, request, send_file,
    redirect, url_for, session
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)

# ---------------------------------------------------------------
# Konfigurasi login & session
# ---------------------------------------------------------------
app.secret_key = os.environ.get("SECRET_KEY", "dev-fallback-ganti-di-production")
APP_PASSWORD = os.environ.get("APP_PASSWORD", "admin123")

app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=7)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"


def login_required(f):
    """Decorator: route hanya bisa diakses kalau sudah login.
    Untuk route /api/*, balikin 401 JSON (bukan redirect) supaya
    frontend gak bingung kalau session-nya expired."""
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            if request.path.startswith("/api/"):
                return jsonify({"error": "Belum login"}), 401
            return redirect(url_for("login", next=request.path))
        return f(*args, **kwargs)
    return decorated


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        if request.form.get("password") == APP_PASSWORD:
            session.permanent = True
            session["logged_in"] = True
            next_url = request.args.get("next") or url_for("home")
            # Cegah open redirect: hanya izinkan path internal
            if not next_url.startswith("/"):
                next_url = url_for("home")
            return redirect(next_url)
        error = "Password salah"
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------------------------------------------------------------
# 1. Load data master
# ---------------------------------------------------------------
DATA_PATH = Path(__file__).parent / "data" / "master_data.json"

with open(DATA_PATH, "r", encoding="utf-8") as f:
    MASTER_DATA = json.load(f)

PREDIKSI_PATH = Path(__file__).parent / "data" / "harga_prediksi.json"

with open(PREDIKSI_PATH, "r", encoding="utf-8") as f:
    PREDIKSI_DATA = json.load(f)


# --- Transform data PIHPS jadi format mirip MASTER_DATA ---
def build_pihps_data():
    """Ubah data historis PIHPS jadi list dengan shape mirip MASTER_DATA,
    biar frontend bisa render ketiganya (Susenas/Seruti/PIHPS) dengan cara
    yang sama. Min/Max = range historis, plus info harga terakhir & bulannya."""
    hasil = []
    for nama, hist in PREDIKSI_DATA["historis_bulanan"].items():
        harga_list = hist["harga"]
        hasil.append({
            "Sumber": "PIHPS",
            "Kode": None,
            "Nama": nama,
            "Satuan": "Kg",
            "Min": min(harga_list),
            "Max": max(harga_list),
            "ButuhInput": None,
            "Label": f"PIHPS - {nama} (Kg)",
            "HargaTerakhir": harga_list[-1],
            "BulanTerakhir": hist["bulan"][-1],
        })
    return hasil


PIHPS_DATA = build_pihps_data()


DATA_BY_LABEL = {item["Label"]: item for item in MASTER_DATA}

KATA_KE_ITEM = {}
for item in MASTER_DATA:
    for kata in item["Nama"].lower().split():
        if len(kata) >= 3:
            KATA_KE_ITEM.setdefault(kata, []).append(item)

def build_sorotan_pasar():
    """Ambil 6 komoditas PIHPS buat sorotan di beranda.
    Setiap entry: nama, harga terakhir, bulan, dan % perubahan dari bulan sebelumnya."""
    daftar = ["Beras", "Cabai Merah", "Cabai Rawit", "Telur Ayam", "Minyak Goreng", "Daging Ayam"]
    hasil = []
    for nama in daftar:
        hist = PREDIKSI_DATA["historis_bulanan"].get(nama)
        if not hist:
            continue
        harga = hist["harga"]
        harga_now = harga[-1]
        harga_prev = harga[-2] if len(harga) >= 2 else harga_now
        persen = ((harga_now - harga_prev) / harga_prev * 100) if harga_prev else 0
        hasil.append({
            "nama": nama,
            "harga": harga_now,
            "bulan": hist["bulan"][-1],
            "persen": round(persen, 1),
            "naik": persen >= 0,
        })
    return hasil


SOROTAN_PASAR = build_sorotan_pasar()

def format_bulan_indo(bulan_str):
    """'2026-08' → 'Agustus 2026'"""
    tahun, bulan = bulan_str.split("-")
    nama = ["Januari","Februari","Maret","April","Mei","Juni",
            "Juli","Agustus","September","Oktober","November","Desember"]
    return f"{nama[int(bulan) - 1]} {tahun}"


# Ambil bulan terakhir dari data historis (dinamis)
_contoh_komoditas = next(iter(PREDIKSI_DATA["historis_bulanan"].values()))
BULAN_TERAKHIR_LABEL = format_bulan_indo(_contoh_komoditas["bulan"][-1])

# ---------------------------------------------------------------
# 1b. Riwayat
# ---------------------------------------------------------------
RIWAYAT_PATH = Path(__file__).parent / "data" / "riwayat.json"


def baca_riwayat():
    if not RIWAYAT_PATH.exists():
        return []
    with open(RIWAYAT_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def catat_riwayat(record):
    data = baca_riwayat()
    data.append(record)
    with open(RIWAYAT_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------
# 2. Flag halaman
# ---------------------------------------------------------------
@app.context_processor
def inject_page_flags():
    return {
        "komoditas_ready": True,
        "about_ready": True,
        "prediksi_ready": True,
    }


# ---------------------------------------------------------------
# 3. Halaman-halaman
# ---------------------------------------------------------------
@app.route("/")
@login_required
def home():
    jumlah_susenas = sum(1 for d in MASTER_DATA if d["Sumber"] == "Susenas")
    jumlah_seruti = sum(1 for d in MASTER_DATA if d["Sumber"] == "Seruti")
    jumlah_pihps = len(PREDIKSI_DATA["historis_bulanan"])
    return render_template(
        "home.html", active_page="home",
        jumlah_komoditas=len(MASTER_DATA),
        jumlah_susenas=jumlah_susenas,
        jumlah_seruti=jumlah_seruti,
        jumlah_pihps=jumlah_pihps,
        sorotan_pasar=SOROTAN_PASAR,
        bulan_terakhir=BULAN_TERAKHIR_LABEL,
    )


@app.route("/cek-kewajaran")
@login_required
def cek_kewajaran():
    return render_template("cek_kewajaran.html", active_page="cek")


@app.route("/komoditas")
@login_required
def komoditas():
    return render_template(
        "komoditas.html", active_page="komoditas",
        data=MASTER_DATA,
        data_pihps=PIHPS_DATA,
    )


@app.route("/prediksi")
@login_required
def prediksi():
    daftar_komoditas = sorted(PREDIKSI_DATA["historis_bulanan"].keys())
    return render_template(
        "prediksi.html", active_page="prediksi",
        daftar_komoditas=daftar_komoditas,
        prediksi_data=PREDIKSI_DATA,
    )


@app.route("/about")
@login_required
def about():
    return render_template("about.html", active_page="about")


def hitung_statistik_riwayat(data):
    hari_ini = datetime.now().date()
    tanggal_7hari = [hari_ini - timedelta(days=i) for i in range(6, -1, -1)]
    hitung_per_tanggal = {t.strftime("%Y-%m-%d"): 0 for t in tanggal_7hari}
    for r in data:
        tgl = r["waktu"].split(" ")[0]
        if tgl in hitung_per_tanggal:
            hitung_per_tanggal[tgl] += 1

    tren_label = [t.strftime("%d %b") for t in tanggal_7hari]
    tren_jumlah = [hitung_per_tanggal[t.strftime("%Y-%m-%d")] for t in tanggal_7hari]

    total_wajar = sum(1 for r in data if r["wajar"])
    total_tidak_wajar = len(data) - total_wajar

    hitung_komoditas = {}
    for r in data:
        kategori = r["nama"].split(" - ")[0]
        hitung_komoditas[kategori] = hitung_komoditas.get(kategori, 0) + 1
    top_komoditas = sorted(hitung_komoditas.items(), key=lambda x: x[1], reverse=True)[:5]

    return {
        "tren_label": tren_label,
        "tren_jumlah": tren_jumlah,
        "total_wajar": total_wajar,
        "total_tidak_wajar": total_tidak_wajar,
        "top_komoditas": top_komoditas,
    }


@app.route("/riwayat")
@login_required
def riwayat():
    data = baca_riwayat()
    statistik = hitung_statistik_riwayat(data)
    data_terbalik = list(reversed(data))
    return render_template(
        "riwayat.html", active_page="riwayat", data=data_terbalik, statistik=statistik
    )


@app.route("/riwayat/hapus", methods=["POST"])
@login_required
def hapus_riwayat():
    with open(RIWAYAT_PATH, "w", encoding="utf-8") as f:
        json.dump([], f, ensure_ascii=False, indent=2)
    return redirect(url_for("riwayat"))


@app.route("/export/riwayat")
@login_required
def export_riwayat():
    data = baca_riwayat()

    dari = request.args.get("dari")
    sampai = request.args.get("sampai")
    if dari:
        data = [r for r in data if r["waktu"][:10] >= dari]
    if sampai:
        data = [r for r in data if r["waktu"][:10] <= sampai]

    wb = Workbook()
    ws = wb.active
    ws.title = "Riwayat"

    header = ["Waktu", "Sumber", "Komoditas", "Satuan", "Harga Diinput (Rp)", "Status"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2A30")
        cell.alignment = Alignment(horizontal="center")

    for r in data:
        ws.append([r["waktu"], r["sumber"], r["nama"], r["satuan"], r["harga"], r["status"]])

    lebar = [20, 10, 40, 10, 18, 30]
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    if dari or sampai:
        rentang = f"{dari or 'awal'}_sd_{sampai or 'akhir'}"
        nama_file = f"riwayat_cek_kewajaran_{rentang}.xlsx"
    else:
        nama_file = f"riwayat_cek_kewajaran_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nama_file,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------
# 4. API
# ---------------------------------------------------------------
@app.route("/api/komoditas")
@login_required
def cari_komoditas():
    query = request.args.get("q", "").strip().lower()
    if not query:
        return jsonify([])

    hasil = [
        item for item in MASTER_DATA
        if query in item["Label"].lower()
    ]

    kata_mirip = difflib.get_close_matches(query, KATA_KE_ITEM.keys(), n=5, cutoff=0.65)

    MAKS_SELISIH_PANJANG = 2
    kata_mirip = [
        kata for kata in kata_mirip
        if abs(len(kata) - len(query)) <= MAKS_SELISIH_PANJANG
    ]

    sudah_ada = {item["Label"] for item in hasil}
    for kata in kata_mirip:
        for item in KATA_KE_ITEM[kata]:
            if item["Label"] not in sudah_ada:
                hasil.append(item)
                sudah_ada.add(item["Label"])

    return jsonify(hasil[:20])


@app.route("/api/cek")
@login_required
def api_cek_harga():
    label = request.args.get("label", "")
    harga_raw = request.args.get("harga", "")
    nama_custom = request.args.get("nama_custom", "").strip()
    satuan_custom = request.args.get("satuan_custom", "").strip()

    item = DATA_BY_LABEL.get(label)
    if item is None:
        return jsonify({"error": "Komoditas tidak ditemukan"}), 404

    try:
        harga = float(harga_raw)
    except ValueError:
        return jsonify({"error": "Harga harus berupa angka"}), 400

    butuh = item.get("ButuhInput")
    if butuh in ("nama", "nama_satuan") and not nama_custom:
        return jsonify({"error": "Nama komoditas spesifik wajib diisi"}), 400
    if butuh == "nama_satuan" and not satuan_custom:
        return jsonify({"error": "Satuan wajib diisi"}), 400

    nama_final = f"{item['Nama']} - {nama_custom}" if nama_custom else item["Nama"]
    satuan_final = satuan_custom if satuan_custom else item["Satuan"]

    is_wajar = item["Min"] <= harga <= item["Max"]
    status = "WAJAR" if is_wajar else "TIDAK WAJAR - CEK ULANG KE RESPONDEN"

    catat_riwayat({
        "waktu": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "nama": nama_final,
        "sumber": item["Sumber"],
        "satuan": satuan_final,
        "harga": harga,
        "status": status,
        "wajar": is_wajar,
    })

    return jsonify({
        "nama": nama_final,
        "sumber": item["Sumber"],
        "satuan": satuan_final,
        "min": item["Min"],
        "max": item["Max"],
        "harga": harga,
        "status": status,
        "wajar": is_wajar,
    })


if __name__ == "__main__":
    app.run(debug=True)